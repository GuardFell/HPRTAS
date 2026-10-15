"""Render `docs/planning/plan-evaluation.md` as a Word document beside it.

The plan and alignment evaluation is kept as Markdown, which is the source of record: it diffs, and
it carries the same identifiers the models, the forms and the tests use. A Word copy is easier to
read and to print, and a second document edited by hand would drift from the first, so this renders
one instead - the same renderer `tests/exports/render-test-plan.py` uses, pointed at the evaluation,
so the two submitted documents read alike. Re-run the script after any change to the evaluation.

    cd docs/planning/exports && python render-plan-evaluation.py

Needs `python-docx`, and the test-plan renderer in the repository (it is imported, not copied, so
there is one Markdown-to-Word renderer and not two).
"""

import importlib.util
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parents[3]
SRC = ROOT / "docs" / "planning" / "plan-evaluation.md"
OUT = ROOT / "docs" / "planning" / "exports" / "plan-evaluation.docx"
RENDERER = ROOT / "tests" / "exports" / "render-test-plan.py"

INTRO = (
    "Plan and alignment evaluation for the Hospital Patient Referral, Treatment and Administration "
    "System: what was built against what was required, the requirement traceability, the stakeholder "
    "goals, and how the plan compares with what happened. Rendered from "
    "`docs/planning/plan-evaluation.md` by `docs/planning/exports/render-plan-evaluation.py`; the "
    "Markdown file is the source of record and this document is a readable copy of it, committed "
    "with it."
)


def load_renderer():
    """The Markdown-to-Word renderer `render-test-plan.py` defines, imported."""
    spec = importlib.util.spec_from_file_location("render_test_plan", RENDERER)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def main():
    renderer = load_renderer()

    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    for attr in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, attr, Cm(1.5))
    doc.styles["Normal"].font.size = Pt(9.5)

    head = doc.add_paragraph()
    renderer.add_runs(head, INTRO)
    for run in head.runs:
        run.italic = True
        run.font.size = Pt(9)
    doc.add_paragraph()

    renderer.render(SRC.read_text(encoding="utf-8").splitlines(), doc)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    renderer.normalise_timestamps(OUT)
    print(f"wrote {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    sys.exit(main())
