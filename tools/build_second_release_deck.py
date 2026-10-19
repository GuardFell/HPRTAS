#!/usr/bin/env python3
"""Build the second-release presentation deck.

Writes `HPRTAS - Second Release.pptx` at the repository root: the deck for the 22 October
2026 submission, in the same visual language as the first-release deck
(`HPRTAS - Five Executable Processes.pptx`) - Segoe UI, navy `#10233A`, slate `#535F6B`,
teal `#1F7A8C`, a navy title slide and white content slides. Every slide's footer names
the artefact its claims come from, and the last slide is the reference list.

The content is the repository at `submission-2026-10-22`: the second release, the
developments that answer the review's feedback, the response record and the plan
evaluation. The numbers are the ones the records of 2026-10-15 name; change a number and
it should change in the documents first.

Run:  python tools/build_second_release_deck.py
Needs: python-pptx
"""

import io
import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(REPO, "HPRTAS - Second Release.pptx")
DIAGRAM_CORE0 = os.path.join(REPO, "models", "exports", "core-0-patient-pathway-end-to-end.png")

FONT = "Segoe UI"
NAVY = RGBColor(0x10, 0x23, 0x3A)
SLATE = RGBColor(0x53, 0x5F, 0x6B)
LIGHT = RGBColor(0xD6, 0xDD, 0xE4)
GREY = RGBColor(0x8A, 0x94, 0x9E)
TEAL = RGBColor(0x1F, 0x7A, 0x8C)
PALE = RGBColor(0xF4, 0xF7, 0xF9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TEALTXT = RGBColor(0x9F, 0xD3, 0xDD)
MUTED = RGBColor(0xC9, 0xD3, 0xDC)
FAINT = RGBColor(0x7E, 0x8C, 0x9A)
GREEN = RGBColor(0x1B, 0x6B, 0x3A)
AMBER = RGBColor(0xA6, 0x6A, 0x00)
RED = RGBColor(0xA3, 0x2A, 0x2A)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
MARGIN = 0.62
WIDTH = 12.1


def textbox(slide, left, top, width, height):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    return frame


def para(frame, text, size=12, color=SLATE, bold=False, first=False, space_after=6, italic=False):
    paragraph = frame.paragraphs[0] if first else frame.add_paragraph()
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.name = FONT
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    paragraph.space_after = Pt(space_after)
    return paragraph


def content_slide(title, footer):
    slide = prs.slides.add_slide(BLANK)
    frame = textbox(slide, MARGIN, 0.30, WIDTH, 0.62)
    para(frame, title, size=22, color=NAVY, bold=True, first=True, space_after=0)
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.64), Inches(0.97),
                                  Inches(12.05), Pt(1.6))
    rule.fill.solid()
    rule.fill.fore_color.rgb = LIGHT
    rule.line.fill.background()
    rule.shadow.inherit = False
    foot = textbox(slide, MARGIN, 7.04, WIDTH, 0.32)
    para(foot, footer, size=8.5, color=GREY, first=True, space_after=0)
    return slide


def bullet_slide(title, footer, items, size=12.5, gap=7):
    slide = content_slide(title, footer)
    frame = textbox(slide, MARGIN, 1.12, WIDTH, 5.8)
    first = True
    for item in items:
        text, level = item[0], item[1]
        opts = item[2] if len(item) > 2 else {}
        prefix = "•  " if level == 0 else "–  "
        paragraph = para(frame, prefix + text,
                         size=opts.get("size", size - (0.5 if level else 0)),
                         color=opts.get("color", SLATE if level == 0 else GREY),
                         bold=opts.get("bold", False),
                         first=first, space_after=opts.get("gap", gap))
        paragraph.level = level
        if level == 0 and not opts.get("bold"):
            for run in paragraph.runs:
                run.font.color.rgb = opts.get("color", NAVY if level == 0 else GREY) \
                    if opts.get("heading") else opts.get("color", SLATE)
        first = False
    return slide


def table_slide(title, footer, data, col_widths, top=1.12, size=9, header_size=9,
                row_height=0.32, note=None, cell_colors=None):
    slide = content_slide(title, footer)
    rows, cols = len(data), len(data[0])
    shape = slide.shapes.add_table(rows, cols, Inches(MARGIN), Inches(top),
                                   Inches(WIDTH), Inches(row_height * rows))
    table = shape.table
    table.first_row = False
    table.horz_banding = False
    total = sum(col_widths)
    for index, width in enumerate(col_widths):
        table.columns[index].width = Inches(WIDTH * width / total)
    for r in range(rows):
        table.rows[r].height = Inches(row_height)
        for c in range(cols):
            cell = table.cell(r, c)
            cell.margin_left = Inches(0.07)
            cell.margin_right = Inches(0.05)
            cell.margin_top = Inches(0.02)
            cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = 1  # middle
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if r == 0 else (PALE if r % 2 == 0 else WHITE)
            frame = cell.text_frame
            frame.word_wrap = True
            paragraph = frame.paragraphs[0]
            run = paragraph.add_run()
            run.text = str(data[r][c])
            run.font.name = FONT
            run.font.size = Pt(header_size if r == 0 else size)
            run.font.bold = r == 0
            default = WHITE if r == 0 else SLATE
            run.font.color.rgb = (cell_colors or {}).get((r, c), default)
            paragraph.space_after = Pt(0)
    if note:
        frame = textbox(slide, MARGIN, top + row_height * rows + 0.16, WIDTH, 0.9)
        para(frame, note, size=10.5, color=SLATE, first=True, space_after=0)
    return slide


# ---------------------------------------------------------------- slide 1, title
slide = prs.slides.add_slide(BLANK)
background = slide.background.fill
background.solid()
background.fore_color.rgb = NAVY
frame = textbox(slide, 0.9, 1.35, 11.5, 5.2)
para(frame, "HPRTAS", size=46, color=WHITE, bold=True, first=True, space_after=4)
para(frame, "Second release  ·  22 October 2026", size=23, color=TEALTXT, space_after=14)
para(frame, "Hospital Patient Referral, Treatment and Administration System",
     size=16, color=MUTED, space_after=10)
para(frame, "The system updated and integrated after the first review: what was built, what the "
            "feedback changed, what is still open and why, how the plan held against what happened, "
            "and how to run and check the whole thing.", size=13, color=MUTED, space_after=18)
para(frame, "github.com/GuardFell/HPRTAS   ·   read it at tag submission-2026-10-22   ·   "
            "the initial release is tag submission-2026-09-29", size=11, color=FAINT, space_after=2)
para(frame, "Every slide's footer names the artefact its claims come from; "
            "slide 16 is the reference and evidence list.", size=11, color=FAINT, space_after=0)

# ---------------------------------------------------------------- slide 2, contents
bullet_slide(
    "What this deck covers",
    "HPRTAS · the deck itself: tools/build_second_release_deck.py; the first-release deck is "
    "HPRTAS - Five Executable Processes.pptx",
    [("The system at a glance — six executable processes, 43 forms, eight workers, a record store", 0),
     ("The second release — what moved between the initial release and the revised version", 0),
     ("Further development — the pathway integrated, the records kept, the forms and the models corrected", 0),
     ("Response to feedback — every point of the 21 review sheets, and what is not done, with the reason", 0),
     ("Project plan evaluation — planned against actual, and where the plan was wrong", 0),
     ("The acceptance criteria and the requirements at the revised version", 0),
     ("Repository, versions and how to run — and where each claim's evidence lives", 0),
     ("References — slide 16 lists the artefacts every slide footer points at", 0)])

# ---------------------------------------------------------------- slide 3, system at a glance
table_slide(
    "The system at a glance",
    "HPRTAS · README.md; models/README.md; forms/README.md; workers/README.md; "
    "tests/evidence/workers-and-models_378b319_2026-10-15.txt",
    [["", "What it is", "Where it lives"],
     ["Processes", "Six executable BPMN processes on Camunda 8 Run 8.9.19: the five pathway "
                   "stages and core-0, which calls them in order", "models/operational/"],
     ["Forms", "43 Camunda Forms; every one of the 61 user tasks and the end-to-end start event "
               "binds one", "forms/"],
     ["Workers", "Eight external workers (Java 21, Maven) and four simulated external services",
      "workers/"],
     ["Records", "The hospital's own store, embedded: patient register, referrals, appointments, "
                 "payment and refund ledger, dispatches, information requests", "workers/…/db/"],
     ["Views", "Two strategic BPMN views and the i* Strategic Dependency and Strategic Rationale "
               "models", "models/strategic/, models/socio-technical/"],
     ["Evidence", "The suite re-run at the revised version: 52 unit tests and 14 engine tests, "
                  "0 failures; 25 of 25 scenarios have a result", "tests/evidence/"]],
    col_widths=[1.3, 7.4, 3.4], row_height=0.72, size=10.5, header_size=10.5)

# ---------------------------------------------------------------- slide 4, what moved
table_slide(
    "Second release: what moved since the initial release",
    "HPRTAS · docs/planning/plan-evaluation.md; tests/test-plan.md; both columns are read at the "
    "tag they name",
    [["Dimension", "Initial release (submission-2026-09-29)", "Revised version (submission-2026-10-22)"],
     ["Processes", "Five independently runnable stages", "The five stages plus core-0, which calls "
      "them as call activities and gates each on what the previous returned"],
     ["Forms", "41", "43 — the reschedule and cancellation journeys and the end-to-end start form"],
     ["Workers", "Seven", "Eight — the document-receipt worker for the documentation exchange"],
     ["Records", "Ledgers in the workers' memory, lost on restart", "The hospital's record store "
      "(embedded H2), read back and used; DEF-10 closed"],
     ["Validation", "Required fields only", "One input-validation policy, applied by a script and "
      "enforced by a check: references, NHS numbers, names, bounds, no past dates"],
     ["Gateways", "Alternative paths merged at tasks", "A named merging gateway at every "
      "convergence - none is left at a task"],
     ["Evidence", "Results recorded against the 29 September tree", "The suite re-run and the two "
      "scenarios that had no result driven at 378b319, 2026-10-15"]],
    col_widths=[1.5, 4.2, 6.4], row_height=0.62, size=10, header_size=10)

# ---------------------------------------------------------------- slide 5, the pathway integrated
slide = content_slide(
    "Further development: the pathway runs as one instance",
    "HPRTAS · models/operational/core-0-patient-pathway-end-to-end.bpmn; "
    "tests/test-plan.md TC-24; tests/evidence/workers-and-models_378b319_2026-10-15.txt")
frame = textbox(slide, MARGIN, 1.10, WIDTH, 3.2)
for text, gap in [
    ("core-0-patient-pathway-end-to-end has no task of its own: five call activities name "
     "core-5, core-1, core-2, core-3 and core-4, and each stage is gated on what the previous "
     "one returned.", 7),
    ("The stages still run on their own — each is started from Tasklist under Processes, so a "
     "single stage or the whole pathway can be demonstrated.", 7),
    ("TC-24 drives one instance through all five stage processes and asserts the tree of "
     "instances the engine recorded: the documentation exchange, the referral and its "
     "appointment, the treatment, the letter and the follow-up, ending at N_CO_PathwayComplete.", 7),
    ("The bindings check grew two rules with this: a call activity has to name a process in "
     "models/operational/, and what a stage returns counts as produced in the process that "
     "calls it.", 7),
]:
    para(frame, "•  " + text, size=12, color=SLATE, first=(gap == 7 and text.startswith("core-0")),
         space_after=gap)
slide.shapes.add_picture(DIAGRAM_CORE0, Inches(0.75), Inches(4.25), width=Inches(11.8))

# ---------------------------------------------------------------- slide 6, the records
bullet_slide(
    "Further development: the hospital's own records",
    "HPRTAS · workers/src/main/java/…/db/Database.java; workers/README.md; "
    "tests/test-plan.md TC-25; tests/evidence/workers-and-models_378b319_2026-10-15.txt",
    [("Every ledger moved out of the workers' memory into an embedded H2 store — the same "
      "database engine Camunda 8 Run keeps its data in: patients, referrals, appointments, "
      "payments, refunds, dispatches, information requests, reference sequence.", 0),
     ("A settled payment is still a duplicate after the store is closed and reopened — the test "
      "closes it and reopens it, which a memory ledger could never show (PersistenceTest; DEF-10 "
      "closed for the part that was a defect).", 0),
     ("The register is read as well as written: a letter that names no channel is sent on the "
      "patient's recorded preference; a rebooking marks the appointment it replaces superseded; "
      "a reply updates the recorded information request.", 0),
     ("Identity is enforced by the worker: an NHS number is ten digits and is the key the register "
      "is looked up by; a name that disagrees with the record is reported, not written over; one "
      "NHS number under a second hospital identifier is a conflict.", 0),
     ("The store is runtime state, not delivered content — workers/data/ is ignored by git, and "
      "mvn compile exec:java -Dexec.args=\"--show-records\" prints what it holds.", 0)],
    size=12.2, gap=10)

# ---------------------------------------------------------------- slide 7, the forms
bullet_slide(
    "Further development: the forms, and what checking them found",
    "HPRTAS · forms/README.md; tools/apply_form_validation.py; tools/check_forms_formjs.mjs; "
    "docs/agile/form-field-burden.md",
    [("One validation policy, not 43 opinions: a reference has to look like one, an NHS number is "
      "ten digits, a field that names a person or a team has to contain a letter, numbers have "
      "bounds, and the dates that are future commitments cannot be picked in the past. "
      "apply_form_validation.py applies it; the bindings check fails if a form drifts.", 0),
     ("Three faults that no run could catch, found by checking: a date field of a type that does "
      "not exist — it rejected whole forms; 26 selects that kept their choices where the library "
      "does not read them, so the picker was empty; groups that nested every field key.", 0),
     ("New journeys with their own ends: reschedule_requested records the change and books again, "
      "marking the old appointment superseded (TC-23); declined_by_patient and cancelled_by_hospital "
      "end at the appointment-cancelled event, not at a confirmation (TC-22).", 0),
     ("@bpmn-io/form-js — the library Tasklist renders with — imports every form and runs each "
      "field through its own validator in a check, in CI-style: an import error is a task that "
      "renders no form at all.", 0),
     ("The review of the recording burden: 43 forms, 293 leaf fields; the clinical lanes ask 8.1 "
      "fields per task on average — counted from the artefacts, with the three findings that ask "
      "for a decision recorded with owners (NFR-009, NFR-011).", 0)],
    size=11.8, gap=8)

# ---------------------------------------------------------------- slide 8, the models
bullet_slide(
    "Further development: the models corrected",
    "HPRTAS · models/operational/; tools/verify_hprtas_bpmn_bindings.py; "
    "tools/verify_hprtas_bpmn_layout.py; tests/test-plan.md section 6",
    [("Merging gateways, named for the convergence they represent — N_OB_MergeContactPath, "
      "N_F_MergeRefundRequest, the N_MS_Merge set and the rest — at every convergence. The two the "
      "first edition left at a task, the payment-error repair paths and the escalation paths, now "
      "merge through N_F_MergePaymentCorrection and N_AM_MergeEscalation.", 0),
     ("Three drawing faults fixed: labels were written over the decision diamonds. The layout now "
      "prefers the routing that keeps lines off the boxes, and the layout check passes on all "
      "eight models; the crossings that remain are the ones the process leaves no other way round.", 0),
     ("DEF-11 closed: an urgent referral with no suitable slot escalates out of the booking process "
      "instead of looping for ever — driven on the engine, the availability check is reached once "
      "and the instance completes (scenario 8).", 0),
     ("DEF-15 closed: an empty optional list is not a problem; an empty required one still is. "
      "Found by driving the reschedule, whose whole point is a second attempt.", 0),
     ("Still open and recorded, not hidden: no user is authenticated (DEF-07) or has submitted a "
      "form (DEF-08); the separation of duties stays structural, and every document says so.", 0)],
    size=11.6, gap=8)

# ---------------------------------------------------------------- slide 9, response 1
table_slide(
    "Response to feedback (1 of 2): the four points that recurred",
    "HPRTAS · docs/agile/response-to-feedback.md — every one of the 21 sheets is accounted for "
    "there, point by point",
    [["Point (sheets that raised it)", "Decision and action", "Evidence", "Outcome"],
     ["The hospital's records — a database, persistence, records used (11)",
      "An embedded H2 store instead of memory; every ledger moved",
      "Database.java; PersistenceTest; TC-25",
      "Duplicates survive a restart; records are read back and used"],
     ["Role-based access — roles, assignment, control of tasks and data (13)",
      "The assignment half: every task carries its lane's candidate group, and the check fails "
      "without one; enforcement recorded as not implemented, not claimed",
      "verify_hprtas_bpmn_bindings.py; FR-045; DEF-07",
      "No task unassigned, and it is said plainly that the lanes are not access control"],
     ["Form validation — required fields, formats, rules, names that are not numbers (15)",
      "One policy, applied by script and enforced by the bindings check",
      "forms/README.md; check_forms_formjs.mjs",
      "A field cannot lose its rule quietly; cross-field rules are stated and refused by the worker"],
     ["Merging gateways where alternative paths come together (12)",
      "Every convergence is a named merging gateway - none left at a task",
      "models/operational/; the layout check",
      "A reader sees where paths join, and the check keeps it true"]],
    col_widths=[3.3, 3.5, 2.5, 3.4], row_height=0.95, size=9, header_size=9.5)

# ---------------------------------------------------------------- slide 10, response 2
table_slide(
    "Response to feedback (2 of 2): the rest, and the weights",
    "HPRTAS · docs/agile/response-to-feedback.md; forms/README.md; "
    "docs/justification-of-decisions.md",
    [["Point (sheets)", "Answer", "Evidence"],
     ["Data types in forms (4)", "Dates are the datetime field's date subtype; selects carry their "
      "choices in values; numbers are number fields", "forms/README.md"],
     ["The model must represent the business process (4)", "The pass's corrections, and four new "
      "check rules — a call activity names a process, a stage's output counts as produced, every "
      "task carries a candidate group, every reference and name carries its pattern",
      "the pass's static-check record"],
     ["Test plan and measurable criteria (3)", "Criteria linked to requirements and business rules; "
      "scenarios with preconditions, data, actions, expected outcome, pass/fail; 25 of 25 now have "
      "a result", "tests/test-plan.md sections 3–5"],
     ["Task types, particularly sending messages (3)", "22 named message flows carry what crosses "
      "a pool boundary; the automated side is a worker against a service task, so no send task "
      "duplicates it", "docs/justification-of-decisions.md section 5"],
     ["Options that have journeys (3)", "Reschedule, cancellation outcomes, the end-to-end start "
      "form, the documentation exchange, the enquiry routing", "TC-20, TC-22, TC-23, TC-24"],
     ["Remove fields or steps (2)", "The burden review measured before changing anything: 293 leaf "
      "fields, clinical lanes 8.1 per task; the three decisions it asks for are recorded with "
      "owners", "docs/agile/form-field-burden.md"],
     ["Dates that cannot be in the past (2)", "disallowPassedDates on the treatment start date, "
      "the rescheduled date and an alternative offered", "forms/README.md"],
     ["Integration (1); NHS identification (1); business logic in workers (1); understanding the "
      "system (2)", "core-0; ten-digit NHS lookup with conflicts reported; eight workers with the "
      "rules enforced where the service is called; the demonstration script as the walkthrough",
      "TC-24, TC-25, workers/README.md, docs/agile/demonstration-script.md"]],
    col_widths=[3.1, 6.0, 3.0], row_height=0.62, size=8.5, header_size=9)

# ---------------------------------------------------------------- slide 11, not done
table_slide(
    "Response to feedback: what is not done, and why",
    "HPRTAS · docs/agile/response-to-feedback.md, section 'Outstanding, and why'; "
    "tests/test-plan.md section 6",
    [["Point", "Why it is not done", "Where it is recorded"],
     ["Enforcement of the roles (the most frequent point — 13 sheets)",
      "Assignment is done and checked; enforcement needs an identity provider and authorisations, "
      "which this prototype has none of. The separation of duties is structural, and every "
      "document says so rather than implying it is enforced", "DEF-07; evaluation section 3"],
     ["A form submitted by a signed-in user",
      "No user is authenticated on the local engine. Every result was driven through the API with "
      "the variables a validated form would supply, and each record states that limit", "DEF-08"],
     ["Dropping fields from referral-review; a translation option in the channel list; a "
      "screen-reader pass",
      "Each changes a variable contract or needs a deployed Tasklist with a user. Proposed with "
      "owners rather than done silently", "docs/agile/form-field-burden.md"],
     ["GP registry validation; multi-instance cardinality; a GP-side rejection",
      "Decisions not to adopt, each with its reason: there is no primary-care registry in the case, "
      "no user task iterates a collection, and the case gives the clinical decision to the hospital",
      "docs/agile/response-to-feedback.md"]],
    col_widths=[3.6, 6.2, 2.3], row_height=1.0, size=9.5, header_size=9.5)

# ---------------------------------------------------------------- slide 12, plan evaluation
table_slide(
    "Project plan evaluation: planned against actual",
    "HPRTAS · docs/planning/project-plan.md section 6; docs/planning/plan-evaluation.md section 5",
    [["Sprint", "The goal as planned", "What actually happened", "Verdict"],
     ["1 — 15–20 Sep", "One deployable model, the acceptance gate, the plan and backlog agreed",
      "Four core-N models, six workers early; the backlog, estimates and contribution tables never "
      "filled in", "Delivery exceeded; the plan process missed"],
     ["2 — 21–27 Sep", "The integrated increment: models, workers and forms end to end",
      "Delivered at a7f0dd6; every service task reached; ten criteria and 21 scenarios defined "
      "ahead of plan", "Ahead of plan"],
     ["3 — 28 Sep–15 Oct", "The i* models and the strategic BPMN, the defects closed, the review's "
      "feedback folded in", "i* models built, checked and exported; DEF-10, DEF-11, DEF-12 and "
      "DEF-15 closed; the feedback pass; two scenarios still unrun inside the sprint",
      "Met, and exceeded; the two closed at the boundary"],
     ["4 — 16–21 Oct", "Validate the revised version; the three evaluations; the demonstration",
      "Suite re-run at 378b319 (52 + 14 tests, 0 failures); TC-13 and TC-16 driven; the criteria, "
      "plan and requirements evaluations written", "On plan; the demonstration is the open row"]],
    col_widths=[1.5, 3.4, 4.6, 2.6], row_height=1.0, size=9.5, header_size=9.5,
    note="Estimates: three sprints carry \"not assessable\" — no estimate was ever agreed, and "
         "back-filling them now would corrupt the velocity figure rather than produce one.")

# ---------------------------------------------------------------- slide 13, where the plan was wrong
bullet_slide(
    "Project plan evaluation: where the plan was wrong",
    "HPRTAS · docs/planning/plan-evaluation.md section 5.4; docs/planning/project-plan.md section 7",
    [("The plan assumed the backlog would fill itself. It cannot be reconstructed afterwards — the "
      "definition of done now requires each task's acceptance conditions to be checked off in the "
      "task's own row, so an unfilled row is visible at the review instead of after it.", 0),
     ("The plan had no task for the model migration. Replacing the first edition with the core-N "
      "models re-pointed every requirement reference by hand and reintroduced two faults that had "
      "already been fixed. A migration is a re-plan, and a re-plan needs its acceptance criteria "
      "re-run.", 0),
     ("A run was treated as evidence of the level above it. The worker test for the refund passed "
      "while the model could not catch the error the worker raised (DEF-12). Evidence has a level, "
      "and a claim made at the wrong level is not evidence.", 0),
     ("Role-based access and the audit trail were left in the requirements as Must and never "
      "planned. They are now explicitly deferred in the plan's scope — a recorded decision, not a "
      "silent omission.", 0),
     ("Estimates were never agreed: three sprints carry \"not assessable\". The cause is not time — "
      "it is that models and workers produce something reviewable and a backlog produces a table "
      "that cannot, so the visible work won.", 0)],
    size=12, gap=11)

# ---------------------------------------------------------------- slide 14, criteria
table_slide(
    "Acceptance criteria at the revised version",
    "HPRTAS · tests/test-plan.md sections 3–6; docs/planning/plan-evaluation.md section 6; "
    "tests/evidence/*_378b319_2026-10-15.txt",
    [["Criterion", "Initial release", "Revised version", "What moved it, and what still limits it"],
     ["AC-01 normal pathway and patient contact", "Met at worker level", "Met",
      "TC-13 drives all three telephone outcomes; the form half remains DEF-08"],
     ["AC-02 clinical decisions and their gates; AC-03 funding and payment; AC-04 exception "
      "routing; AC-05 no suitable slot and urgent referrals; AC-08 refund and separation; AC-09 "
      "data minimisation; AC-10 controlled failure",
      "Met", "Met",
      "Re-run at 378b319; all three referral decisions are now driven, where TC-02 had two"],
     ["AC-06 clinic letter, seven-day target and escalation", "Partially met — the timer's path "
      "had never been run", "Partially met, on a better basis: the timer fires with no worker, "
      "and the reminder and escalation are driven (TC-16)",
      "What remains is that the escalation thresholds are recorded data, not measured periods — "
      "said in the evaluation rather than rounded up"],
     ["AC-07 follow-up, cancellation and enquiries", "Partially met", "Partially met",
      "The routing is driven; the record BR-37 requires still comes from the API, not a signed-in "
      "user (DEF-08)"]],
    col_widths=[3.4, 2.2, 3.3, 3.2], row_height=0.95, size=9, header_size=9.5,
    note="Every scenario now has a result: 25 of 25 — 18 pass and 7 pass in part. The criteria "
         "text was not rewritten; the second release revises the verdicts and the runs behind them.")

# ---------------------------------------------------------------- slide 15, repository and running
slide = content_slide(
    "Repository, versions, and how to run it",
    "HPRTAS · README.md 'Running it'; tools/README.md; models/operational/README.md")
frame = textbox(slide, MARGIN, 1.10, 5.85, 5.6)
para(frame, "Versions", size=13, color=NAVY, bold=True, first=True, space_after=5)
for text in ["Read the revised version at tag submission-2026-10-22.",
             "The initial release is tag submission-2026-09-29 (release-2.0 names the same tree); "
             "release-1.0 is the first tagged release and is left where it is, because the results "
             "recorded against it are only readable there.",
             "Every result in tests/evidence/ names the commit it was produced at."]:
    para(frame, "•  " + text, size=11.5, color=SLATE, space_after=6)
para(frame, "Run it", size=13, color=NAVY, bold=True, space_after=5)
for text in ["camunda-runtime\\start-camunda.bat — the engine (Java 21; the first start takes "
             "about a minute).",
             "Deploy models/operational/*.bpmn and forms/*.form together — a task resolves its form "
             "by key.",
             "Workers: cd workers && mvn compile exec:java; the records land in workers/data/ "
             "(HPRTAS_DB_URL moves the store).",
             "Tasklist at localhost:8080/tasklist, login demo / demo."]:
    para(frame, "•  " + text, size=11.5, color=SLATE, space_after=6)
frame = textbox(slide, 6.85, 1.10, 5.85, 5.6)
para(frame, "Check it", size=13, color=NAVY, bold=True, first=True, space_after=5)
for text in ["python tools/verify_hprtas_bpmn_bindings.py — the models, forms and workers agree.",
             "python tools/verify_hprtas_bpmn_layout.py — the layout and the BPMN rules.",
             "python tools/verify_hprtas_istar.py and verify_hprtas_pistar.py — the i* models and "
             "their piStar files.",
             "cd workers && mvn test, then mvn test -Pengine — 52 unit and 14 engine tests.",
             "The checks need no engine except the engine-form one, and each says what it did "
             "not cover."]:
    para(frame, "•  " + text, size=11.5, color=SLATE, space_after=6)
para(frame, "The written evaluations", size=13, color=NAVY, bold=True, space_after=5)
for text in ["docs/planning/exports/plan-evaluation.docx and tests/exports/test-plan.docx are the "
             "Word renderings of the two evaluation documents."]:
    para(frame, "•  " + text, size=11.5, color=SLATE, space_after=6)

# ---------------------------------------------------------------- slide 16, references
bullet_slide(
    "References and evidence links",
    "HPRTAS · every path below is in the repository; read it at tag submission-2026-10-22",
    [("Repository: github.com/GuardFell/HPRTAS — the revised version is tag submission-2026-10-22; "
      "the initial release is submission-2026-09-29.", 0),
     ("The evaluation: docs/planning/plan-evaluation.md — requirement traceability (section 3), "
      "the models against the forms and workers (4), the plan (5), the criteria (6); Word form at "
      "docs/planning/exports/plan-evaluation.docx.", 0),
     ("The requirements and their traceability: docs/requirements/requirements.md.", 0),
     ("The tests: tests/test-plan.md — criteria (3), scenarios (4), the execution record (5), "
      "defects and limitations (6) — and the runs in tests/evidence/, including the four records "
      "of 2026-10-15.", 0),
     ("The response to the feedback: docs/agile/response-to-feedback.md.", 0),
     ("The demonstration: docs/agile/demonstration-script.md — the flow, the commit and the limits.", 0),
     ("The models: models/README.md; models/strategic/; models/socio-technical/ (i* SD and SR, "
      "with the piStar files); models/operational/; the drawings in models/exports/ as PDF and PNG.", 0),
     ("The workers and the forms: workers/README.md, forms/README.md; the checks and generators: "
      "tools/README.md.", 0),
     ("The group's own record: Group Contribution Table.docx, at the repository root.", 0)],
    size=11.5, gap=8)

prs.save(OUT)
print(f"wrote {OUT} with {len(prs.slides._sldIdLst)} slides")
